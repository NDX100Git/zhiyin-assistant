"""文档解析（阶段 5-A）测试：PDF 解析、空白页、页边距、图片、参考文献、缓存。"""

from __future__ import annotations

import io
import shutil
import zipfile
from pathlib import Path

import fitz
import pytest
from _cjk_font import cjk_font_args
from sqlalchemy import func, select

from app.config import settings
from app.models.ai_report import AiReport
from app.services.doc_parser import (
    detect_sections,
    extract_references,
    parse_document,
    parse_pdf,
    quick_page_count,
)
from app.services.storage import build_object_key


def _write_text(page: fitz.Page, point: tuple[float, float], text: str, size: int = 12) -> None:
    page.insert_text(point, text, fontsize=size, **cjk_font_args())


def _make_sample_pdf(path: Path) -> None:
    """构造 4 页样例：摘要（含图片）、空白页、正文、参考文献。"""
    png_path = path.parent / "block.png"
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 120, 80))
    pix.clear_with(200)
    pix.save(str(png_path))

    with fitz.open() as doc:
        page1 = doc.new_page()
        _write_text(page1, (90, 100), "摘要", 16)
        _write_text(page1, (90, 130), "本文研究了论文格式自动检查方法，提出了规则与大模型混合的架构。", 12)
        page1.insert_image(fitz.Rect(320, 200, 460, 300), filename=str(png_path))

        doc.new_page()  # 第 2 页：空白页

        page3 = doc.new_page()
        _write_text(page3, (90, 100), "第一章 绪论", 16)
        _write_text(page3, (90, 130), "本文的第一章介绍研究背景与意义。", 12)

        page4 = doc.new_page()
        _write_text(page4, (90, 100), "参考文献", 16)
        _write_text(page4, (90, 130), "[1] 张三. 论文题目[J]. 某期刊, 2020, 12(3): 45-52.", 12)
        _write_text(page4, (90, 150), "[2] 李四. 专著名称[M]. 北京: 科学出版社, 2019.", 12)

        doc.save(str(path))


@pytest.fixture
def sample_pdf(tmp_path) -> Path:
    path = tmp_path / "sample.pdf"
    _make_sample_pdf(path)
    return path


def test_解析PDF得到页数文本与空白页(sample_pdf):
    report = parse_pdf(sample_pdf)

    assert report["page_count"] == 4
    assert report["blank_pages"] == [2]
    assert report["pages_with_images"] == [1]
    assert report["text_chars"] > 0
    assert report["pages"][0]["char_count"] > 0
    assert report["pages"][1]["blank"] is True


def test_解析结果包含排版属性(sample_pdf):
    from app.services.doc_parser import normalize_font_name

    report = parse_pdf(sample_pdf)
    typography = report["typography"]

    assert "available" in typography
    assert isinstance(typography.get("font_histogram", []), list)
    # 字体名归一化：子集前缀与英文别名都要能映射
    assert normalize_font_name("ABCDEF+SimSun") == "宋体"
    assert normalize_font_name("SimHei-Bold") == "黑体"
    assert normalize_font_name("思源宋体") == "思源宋体"


def test_页边距与图片位置以厘米返回(sample_pdf):
    report = parse_pdf(sample_pdf)
    first = report["pages"][0]

    margins = first["margins_cm"]
    assert margins is not None
    assert 0 < margins["top_cm"] < 10
    assert 0 < margins["left_cm"] < 10

    assert first["image_count"] == 1
    image = first["images"][0]
    assert image["width_cm"] > 0 and image["height_cm"] > 0
    assert image["x_cm"] > 0
    # 图片像素宽高与按打印尺寸估算的有效分辨率
    assert image["pixel_width"] == 120
    assert image["pixel_height"] == 80
    assert image["dpi"] is not None and image["dpi"] < 150


def test_识别章节与参考文献条目(sample_pdf):
    report = parse_pdf(sample_pdf)

    assert "摘要" in report["sections"]
    assert "参考文献" in report["sections"]
    assert "正文" in report["sections"]  # 由“第一章 绪论”识别
    assert report["sections"].index("摘要") < report["sections"].index("参考文献")

    references = report["references"]
    assert [item["index"] for item in references] == [1, 2]
    assert "45-52" in references[0]["text"]


def test_章节与参考文献解析函数():
    text = "摘要\n本文\n第一章 绪论\n内容\n参考文献\n[1] 王五. 标题[J]. 期刊, 2021, 3(1): 1-9.\n"
    assert detect_sections(text) == ["摘要", "正文", "参考文献"]
    assert len(extract_references(text)) == 1


def test_docx页数快读(tmp_path):
    with_pages = tmp_path / "with_pages.docx"
    with zipfile.ZipFile(with_pages, "w") as zf:
        zf.writestr(
            "docProps/app.xml",
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties">'
            "<Pages>7</Pages><Words>1234</Words></Properties>",
        )
    assert quick_page_count(with_pages) == 7

    without_pages = tmp_path / "without_pages.docx"
    with zipfile.ZipFile(without_pages, "w") as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
    assert quick_page_count(without_pages) is None

    not_docx = tmp_path / "broken.docx"
    not_docx.write_bytes(b"not a zip")
    assert quick_page_count(not_docx) is None


def test_解析结果写入缓存且不重复解析(db_session, test_storage, sample_pdf):
    key = build_object_key("sample.pdf")
    data = sample_pdf.read_bytes()
    test_storage.upload(key, io.BytesIO(data), len(data), "application/pdf")

    first = parse_document(db_session, key)
    second = parse_document(db_session, key)

    assert first["page_count"] == second["page_count"] == 4
    assert second["file_key"] == key
    cached_rows = db_session.execute(select(func.count()).select_from(AiReport)).scalar_one()
    assert cached_rows == 1


@pytest.mark.skipif(
    settings.soffice_binary is None,
    reason="本机未配置可用的 LibreOffice，跳过 Office 转换测试",
)
def test_docx经LibreOffice转换后可解析(db_session, test_storage, tmp_path):
    from docx import Document

    docx_path = tmp_path / "thesis.docx"
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph("本文研究论文格式检查方法。" * 20)
    document.add_page_break()
    document.add_heading("参考文献", level=1)
    document.add_paragraph("[1] 张三. 论文题目[J]. 某期刊, 2020, 12(3): 45-52.")
    document.save(str(docx_path))

    key = build_object_key("thesis.docx")
    data = docx_path.read_bytes()
    test_storage.upload(key, io.BytesIO(data), len(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    report = parse_document(db_session, key)
    assert report["source"]["converted_from"] == "docx"
    assert report["page_count"] >= 1
    assert "摘要" in "".join(page["text"] for page in report["pages"])


@pytest.mark.skipif(
    not settings.soffice_container or shutil.which("docker") is None,
    reason="未配置可用的 LibreOffice 转换容器，跳过容器内 PPT 转换测试",
)
def test_pptx经转换容器转换后可解析(db_session, test_storage, tmp_path):
    """容器镜像必须带 impress，否则上传接口声称支持的 PPT/PPTX 会在部署后失败。"""
    from pptx import Presentation

    pptx_path = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "论文答辩"
    slide.placeholders[1].text = "规则引擎与模型混合的格式检查"
    presentation.save(str(pptx_path))

    key = build_object_key("slides.pptx")
    data = pptx_path.read_bytes()
    test_storage.upload(
        key,
        io.BytesIO(data),
        len(data),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )

    report = parse_document(db_session, key)
    assert report["source"]["converted_from"] == "pptx"
    assert report["page_count"] >= 1


# ---------------------------------------------------------------- 章节定位
def test_正文里提到参考文献不误判为章节标题():
    """「本文的参考文献均按 GB/T 7714 著录」不是标题。

    早期实现用全文 find：正文里出现这四个字就被当成章节，而且只取第一个匹配，
    真正的标题反而看不到 —— 章节顺序、结构完整性、文献条目提取全被带偏。
    """
    text = (
        "摘要\n本文研究论文格式自动检查方法。\n"
        "本文的参考文献均按 GB/T 7714 著录，具体条目见文末。\n"
        "第一章 绪论\n参考文献的排版也会影响查重结果。\n"
        "参考文献\n[1] 张三. 题目[J]. 期刊, 2020, 12(3): 45-52.\n"
    )

    assert detect_sections(text) == ["摘要", "正文", "参考文献"]

    items = extract_references(text)
    assert [item["index"] for item in items] == [1]
    assert "45-52" in items[0]["text"]


def test_参考文献标题带冒号或英文时仍能识别():
    for heading in ("参考文献：", "参考文献 References", "参考文献"):
        text = f"摘要\n正文内容\n{heading}\n[1] 王五. 标题[J]. 期刊, 2021, 3(1): 1-9.\n"
        assert "参考文献" in detect_sections(text), heading
        assert [item["index"] for item in extract_references(text)] == [1], heading


def test_正文里的编号列表不会被当成参考文献条目():
    """正文里提到「参考文献」后紧跟 [1] 编号列表时，不能把正文列表当文献著录。"""
    text = (
        "摘要\n本文研究格式检查。\n"
        "参考文献的取舍标准如下：[1] 只收公开发表的成果。\n"
        "第一章 绪论\n内容\n"
        "参考文献\n[1] 张三. 题目[J]. 期刊, 2020, 12(3): 45-52.\n"
    )

    items = extract_references(text)
    assert len(items) == 1
    assert items[0]["text"].startswith("张三")


# ---------------------------------------------------------------- 转换子进程的环境变量
def test_转换子进程只拿到最小环境变量(monkeypatch, tmp_path):
    """后端环境里有数据库口令、JWT 密钥、模型 API Key，不能传给 LibreOffice。

    LibreOffice 会派生一串子进程，转换的又是学生上传的任意文档；宏、扩展、
    崩溃报告任何一处把环境打出来，密钥就跟着泄露。
    """
    from app.services import doc_parser

    monkeypatch.setenv("DB_PASSWORD", "should-not-leak-db")
    monkeypatch.setenv("JWT_SECRET", "should-not-leak-jwt")
    monkeypatch.setenv("LLM_API_KEY", "sk-should-not-leak-llm")

    out_dir = tmp_path / "work"
    out_dir.mkdir()
    env = doc_parser._child_env(out_dir)

    assert "DB_PASSWORD" not in env
    assert "JWT_SECRET" not in env
    assert "LLM_API_KEY" not in env
    leaked = " ".join(env.values())
    for secret in ("should-not-leak-db", "should-not-leak-jwt", "sk-should-not-leak-llm"):
        assert secret not in leaked

    # 跑起来需要的变量还在
    assert env.get("PATH")
    assert env["TMP"] == env["TEMP"]
    assert Path(env["TMP"]).is_dir()


def test_启动转换进程时传的是白名单环境(monkeypatch, tmp_path):
    from app.services import doc_parser

    captured: dict = {}

    class _FakeProc:
        pid = 4242

        def wait(self, timeout=None):
            return 0

    def fake_popen(cmd, **kwargs):
        captured.update(kwargs)
        return _FakeProc()

    monkeypatch.setattr(doc_parser.subprocess, "Popen", fake_popen)
    monkeypatch.setenv("DB_PASSWORD", "should-not-leak-db")

    out_dir = tmp_path / "w"
    out_dir.mkdir()
    doc_parser._run_soffice(["true"], out_dir, out_dir / "o.log", out_dir / "e.log")

    assert "DB_PASSWORD" not in captured["env"]
    assert "should-not-leak-db" not in " ".join(captured["env"].values())
    assert captured["env"]["TMP"] == captured["env"]["TEMP"]
