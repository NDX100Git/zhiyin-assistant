"""表格归一版式的测试：不依赖 LibreOffice，纯文件级校验。

这些断言盯的是「页数确定」的前提：打印区域收紧、缩到一页宽、分页符清干净、
隐藏表不出纸。任何一条回退，同一个表格的页数就会随排版漂移。
"""

from __future__ import annotations

import time
from pathlib import Path

import pymupdf
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.pagebreak import Break

from app.config import settings
from app.services import spreadsheet
from app.services.spreadsheet import normalize_spreadsheet, render_pdf
from app.utils.exceptions import BusinessError


def _wide_sheet(path: Path, rows: int = 30, cols: int = 12) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "汇总"
    for row in range(1, rows + 1):
        for col in range(1, cols + 1):
            sheet.cell(row=row, column=col, value=f"R{row}C{col}")
    # 尾部留出空行空列，模拟学生表里常见的「幽灵区」。
    # 只赋值 None 不会被写入文件，要带格式才会在加载后撑大 max_row/max_column
    ghost = sheet.cell(row=rows + 20, column=cols + 10)
    ghost.value = None
    ghost.font = Font(bold=True)
    sheet.row_breaks.append(Break(id=10))
    workbook.save(path)
    return path


def test_表格归一化把打印区域收紧到有内容的行列(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "汇总表.xlsx", rows=30, cols=12)
    target = normalize_spreadsheet(source, tmp_path / "work")

    sheet = load_workbook(target).active
    # openpyxl 会把表名一起写进打印区域，这里只校验单元格范围
    assert str(sheet.print_area).endswith("$A$1:$L$30")


def test_表格归一化打开一页宽并清掉手动分页符(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "宽表.xlsx")
    target = normalize_spreadsheet(source, tmp_path / "work")

    sheet = load_workbook(target).active
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert len(sheet.row_breaks.brk) == 0


def test_隐藏工作表不参与打印(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active.title = "要打印"
    workbook.active["A1"] = "明细"
    workbook.active["B2"] = "第二格"
    hidden = workbook.create_sheet("隐藏明细")
    hidden["A1"] = "不该出纸"
    hidden.sheet_state = "hidden"
    source = tmp_path / "多表.xlsx"
    workbook.save(source)

    target = normalize_spreadsheet(source, tmp_path / "work")
    result = load_workbook(target)
    assert result["隐藏明细"].sheet_state == "hidden"
    assert str(result["要打印"].print_area).endswith("$A$1:$B$2")


def test_看不见内容的工作表被跳过(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = "有内容"
    empty = workbook.create_sheet("空白表")
    empty["C5"] = "   "
    source = tmp_path / "含空表.xlsx"
    workbook.save(source)

    target = normalize_spreadsheet(source, tmp_path / "work")
    result = load_workbook(target)
    assert result["空白表"].sheet_state == "hidden"


def test_整张表都空时明确报错(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.active["A1"] = None
    source = tmp_path / "空表.xlsx"
    workbook.save(source)

    with pytest.raises(BusinessError) as excinfo:
        normalize_spreadsheet(source, tmp_path / "work")
    assert "没有任何内容" in str(excinfo.value)


def test_损坏文件给出人话提示(tmp_path: Path) -> None:
    source = tmp_path / "坏表.xlsx"
    source.write_bytes(b"not an excel file")

    with pytest.raises(BusinessError) as excinfo:
        normalize_spreadsheet(source, tmp_path / "work")
    assert "打不开" in str(excinfo.value)


def test_横向分成两页时首列作为标签列重复(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "宽表.xlsx")
    target = normalize_spreadsheet(source, tmp_path / "work", fit_to_width=2)

    sheet = load_workbook(target).active
    assert sheet.page_setup.fitToWidth == 2
    assert sheet.print_title_cols == "$A:$A"


def _fake_convert(font_pt: float, calls: list[int]):
    """替身转换器：产出一份指定字号的 PDF，省掉真跑 LibreOffice。"""

    def convert(prepared: Path, out_dir: Path) -> Path:
        calls.append(1)
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "表格样例", fontsize=font_pt)
        target = out_dir / f"{prepared.stem}.pdf"
        document.save(target)
        document.close()
        return target

    return convert


def test_字号够大就停在一页宽(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "正常表.xlsx")
    calls: list[int] = []
    stats: dict = {}

    render_pdf(source, tmp_path / "work", _fake_convert(11.0, calls), stats=stats)
    assert len(calls) == 1
    # 没有降级就不该写降级说明，否则返回体上会挂一句误导人的话
    assert "render_note" not in stats
    assert "render_truncated" not in stats


def test_字号太小自动放宽横向页数(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "超宽表.xlsx")
    calls: list[int] = []

    pdf = render_pdf(source, tmp_path / "work", _fake_convert(5.9, calls))
    # 一直小：放宽到上限（默认 4 页宽）后仍交出一份 PDF，不抛错
    assert len(calls) == 4
    assert pdf.exists()


def test_归一化带出每个表的末尾空行空列(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "带空白区.xlsx", rows=30, cols=12)
    stats: dict = {}

    normalize_spreadsheet(source, tmp_path / "work", stats=stats)

    sheet = stats["sheets"][0]
    assert sheet["title"] == "汇总"
    assert sheet["content_rows"] == 30
    assert sheet["content_columns"] == 12
    # _wide_sheet 在 50 行、22 列留了幽灵单元格
    assert sheet["trailing_blank_rows"] == 20
    assert sheet["trailing_blank_cols"] == 10


def test_渲染结果带出实际中位字号与横向页数(tmp_path: Path) -> None:
    source = _wide_sheet(tmp_path / "可读表.xlsx")
    stats: dict = {}

    render_pdf(source, tmp_path / "work", _fake_convert(11.0, []), stats=stats)

    assert stats["fit_to_width_pages"] == 1
    assert stats["print_font_pt"] == 11.0
    assert stats["min_readable_font_pt"] > 0
    assert stats["sheets"][0]["trailing_blank_rows"] == 20
    # 实测数字随结果一起带出，验收时要看得到次数与耗时
    assert stats["render_attempts"] == 1
    assert stats["render_elapsed_seconds"] >= 0


# ---------------------------------------------------------------- 限次与预算
# 验收实测过的反例：「每次转换 130 秒时仍跑 2 次、占用 260 秒」，而排在后面的
# 请求 20 秒就被挤成超时。所以次数与总耗时都必须是硬约束，且额度要传到真正等
# soffice 的那一层（doc_parser.conversion_timeout_seconds），光在循环里数次数
# 挡不住「下一轮自己就能跑满 180 秒」。
def _scripted_convert(each: float, calls: list[float], font_pt: float = 5.9):
    """替身转换器：每轮耗时 each 秒；额度不够就当场超时，仿真杀进程的 50041。"""

    def convert(prepared: Path, out_dir: Path) -> Path:
        remaining = spreadsheet.remaining_render_seconds()
        wait = each if remaining is None else min(each, remaining)
        calls.append(wait)
        time.sleep(wait)
        if remaining is not None and wait < each:
            raise BusinessError("文档转换超时，请稍后重试", code=50041)
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), "表格样例", fontsize=font_pt)
        target = out_dir / f"{prepared.stem}.pdf"
        document.save(target)
        document.close()
        return target

    return convert


def test_慢转换被预算截住并按前一次结果出纸(monkeypatch, tmp_path: Path) -> None:
    """每轮 0.5 秒、预算 0.9 秒：第一轮跑满，第二轮只剩 0.4 秒，到点就停。

    按真实比例缩放（240 秒预算 / 130 秒一轮）。旧实现里第二轮不受剩余时间约束，
    能跑满自己的一次转换超时，两轮加起来就超预算。
    """
    source = _wide_sheet(tmp_path / "慢表.xlsx")
    calls: list[float] = []
    stats: dict = {}
    monkeypatch.setattr(settings, "spreadsheet_render_budget_seconds", 0.9)
    monkeypatch.setattr(settings, "soffice_timeout_seconds", 180)

    started = time.monotonic()
    pdf = render_pdf(source, tmp_path / "work", _scripted_convert(0.5, calls), stats=stats)
    elapsed = time.monotonic() - started

    max_width = int(spreadsheet.get_spreadsheet_rules()["max_width_pages"])
    assert pdf.exists()
    assert len(calls) <= max_width, "转换次数不能超过配置里的最大页宽"
    assert len(calls) == 2, "第一轮跑完还有余额，第二轮被额度卡死"
    assert sum(calls) <= 0.9 + 0.02, f"转换总耗时要被预算截住，实际 {sum(calls):.3f} 秒"
    assert elapsed <= 0.9 + 0.3, f"整条请求的耗时要被预算截住，实际 {elapsed:.3f} 秒"

    assert stats["fit_to_width_pages"] == 1, "降级后按第一轮的版式出纸"
    assert stats["render_attempts"] == 1, "第二轮没拿到完整结果，不能计入已用轮次"
    assert stats["render_truncated"] is True
    assert "因超时按前一次结果出纸" in stats["render_note"]


def test_预算一次用尽时不再起第二轮(monkeypatch, tmp_path: Path) -> None:
    """第一轮就把预算跑满时，直接按这一轮出纸，不再试第二轮。"""
    source = _wide_sheet(tmp_path / "刚够表.xlsx")
    calls: list[float] = []
    stats: dict = {}
    monkeypatch.setattr(settings, "spreadsheet_render_budget_seconds", 0.2)
    monkeypatch.setattr(settings, "soffice_timeout_seconds", 180)

    pdf = render_pdf(source, tmp_path / "work", _scripted_convert(0.3, calls), stats=stats)

    assert pdf.exists()
    assert len(calls) == 1, "第一轮已用掉 0.3 秒、预算只有 0.2 秒，不该再起一轮"
    assert stats["render_truncated"] is True
    assert "因超时按前一次结果出纸" in stats["render_note"]


def test_渲染预算传到了转换层(monkeypatch, tmp_path: Path) -> None:
    """预算必须传到真正等 soffice 的那一层，否则一轮照样能等满 180 秒。"""
    source = _wide_sheet(tmp_path / "预算表.xlsx")
    seen: list[float | None] = []
    instant = _fake_convert(11.0, [])

    def convert(prepared: Path, out_dir: Path) -> Path:
        seen.append(spreadsheet.remaining_render_seconds())
        return instant(prepared, out_dir)

    assert spreadsheet.remaining_render_seconds() is None, "不在表格渲染里时不受约束"
    monkeypatch.setattr(settings, "spreadsheet_render_budget_seconds", 0.5)
    monkeypatch.setattr(settings, "soffice_timeout_seconds", 180)

    render_pdf(source, tmp_path / "work", convert)

    assert seen and seen[0] is not None
    assert seen[0] > 0.4, "第一轮拿到的是预算与单次超时里更大的那个额度"
    assert spreadsheet.remaining_render_seconds() is None, "出了 render_pdf 就该复位"
