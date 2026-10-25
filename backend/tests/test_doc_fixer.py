"""文档修复（阶段 6）测试：指令校验、指令执行、修复副本、PDF 先转 Word。"""

from __future__ import annotations

import io

import fitz
import pytest
from _cjk_font import cjk_font_args
from docx import Document

from app.services.doc_fixer import apply_actions, fix_document, sanitize_actions
from app.services.doc_parser import save_report
from app.services.storage import build_object_key
from app.utils.exceptions import BusinessError


def _make_docx(path, reference_text: str) -> None:
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph("本文研究论文格式检查方法，实验结杲表明方法有效。")
    document.add_page_break()             # 分页符
    document.add_paragraph("")            # 空白页上的空段落
    document.add_paragraph("")
    document.add_paragraph("参考文献")
    document.add_paragraph(reference_text)
    document.save(str(path))


def _check_report(file_key: str, reference_text: str) -> dict:
    return {
        "file_key": file_key,
        "issues": [
            {
                "id": "i1",
                "type": "structure",
                "severity": "high",
                "location": "全文",
                "original": "未检测到「目录」",
                "suggestion": "补充目录",
                "reason": "GB/T 7713.1",
                "source": "rule",
            },
            {
                "id": "i2",
                "type": "reference",
                "severity": "high",
                "location": "参考文献 第 1 条",
                "original": reference_text,
                "suggestion": "补齐卷(期):页码",
                "reason": "GB/T 7714",
                "source": "ai",
            },
            {
                "id": "i3",
                "type": "layout",
                "severity": "medium",
                "location": "第 2 页",
                "original": "该页文本极少且没有图片",
                "suggestion": "删除空白页",
                "reason": "浪费纸张",
                "source": "rule",
            },
        ],
    }


def test_指令校验丢弃非法项():
    raw = {
        "actions": [
            {"op": "set_line_spacing", "value": 1.5},
            {"op": "set_line_spacing", "value": 9},
            {"op": "set_font", "value": {"cn": "宋体", "size": 12}},
            {"op": "set_font", "value": {"cn": "", "size": 12}},
            {"op": "remove_blank_page"},
            {"op": "fix_reference", "original": "旧条目", "value": "新条目"},
            {"op": "delete_file", "value": "x"},
            {"op": "fix_reference", "original": "旧条目", "value": ""},
        ]
    }
    valid, skipped = sanitize_actions(raw)

    assert [item["op"] for item in valid] == [
        "set_line_spacing",
        "set_font",
        "remove_blank_page",
        "fix_reference",
    ]
    assert len(skipped) == 4
    assert any(item["reason"] == "不支持的指令类型" for item in skipped)


def test_修复docx生成副本且指令生效(db_session, test_storage, tmp_path, monkeypatch):
    reference_text = "[1] 张三. 论文格式自动检查方法研究[J]. 计算机应用, 2020, 12(3)."
    source = tmp_path / "source.docx"
    _make_docx(source, reference_text)
    key = build_object_key("source.docx")
    data = source.read_bytes()
    test_storage.upload(key, io.BytesIO(data), len(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    save_report(db_session, key, _check_report(key, reference_text), "check")

    fixed_reference = (
        "[1] 张三. 论文格式自动检查方法研究[J]. 计算机应用, 2020, 12(3): 45-52."
    )
    monkeypatch.setattr(
        "app.services.doc_fixer.chat_json",
        lambda messages: {
            "actions": [
                {"op": "set_line_spacing", "target": "body", "value": 1.5},
                {"op": "set_font", "target": "body", "value": {"cn": "宋体", "size": 12}},
                {"op": "remove_blank_page", "target": "blank_pages"},
                {
                    "op": "fix_reference",
                    "target": "参考文献 第 1 条",
                    "original": reference_text,
                    # 故意不带 [1] 序号，验证代码会把序号补回来
                    "value": fixed_reference.replace("[1] ", ""),
                },
            ]
        },
    )

    result = fix_document(db_session, key, ["i2", "i3"], file_name="修复结果.docx")

    assert result["source_file_key"] == key
    assert result["file_key"] != key
    assert result["converted_from_pdf"] is False
    assert result["selected_issue_ids"] == ["i2", "i3"]
    assert len(result["applied_actions"]) == 4
    assert result["failed_actions"] == []

    # 下载修复后的副本并验证
    fixed_bytes = test_storage.read_bytes(result["file_key"])
    fixed_doc = Document(io.BytesIO(fixed_bytes))
    texts = [paragraph.text for paragraph in fixed_doc.paragraphs]

    assert any(fixed_reference in text for text in texts)
    assert not any(reference_text == text for text in texts)
    # 分页符后面的两段空段落才是空白页，被清掉；分页符本身保留
    assert sum(1 for text in texts if not text.strip()) <= 1
    body = next(paragraph for paragraph in fixed_doc.paragraphs if "本文研究" in paragraph.text)
    assert body.paragraph_format.line_spacing == 1.5
    assert body.runs[0].font.name == "宋体"

    # 原文件保持不变
    original = Document(io.BytesIO(test_storage.read_bytes(key)))
    assert any(reference_text in paragraph.text for paragraph in original.paragraphs)


def test_未勾选问题时报错(db_session, test_storage, tmp_path):
    source = tmp_path / "source.docx"
    _make_docx(source, "[1] 张三. 题目[J]. 期刊, 2020.")
    key = build_object_key("source.docx")
    data = source.read_bytes()
    test_storage.upload(key, io.BytesIO(data), len(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    save_report(db_session, key, _check_report(key, "[1] 张三. 题目[J]. 期刊, 2020."), "check")

    with pytest.raises(BusinessError) as exc:
        fix_document(db_session, key, [])
    assert "勾选" in exc.value.message

    with pytest.raises(BusinessError) as exc2:
        fix_document(db_session, key, ["not-exist"])
    assert "没有找到" in exc2.value.message


def test_没有体检报告时提示先体检(db_session, test_storage, tmp_path):
    source = tmp_path / "source.docx"
    _make_docx(source, "[1] 张三. 题目[J]. 期刊, 2020.")
    key = build_object_key("source.docx")
    data = source.read_bytes()
    test_storage.upload(key, io.BytesIO(data), len(data), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    with pytest.raises(BusinessError) as exc:
        fix_document(db_session, key, ["i1"])
    assert "体检" in exc.value.message


def test_PDF自动转Word后修复(db_session, test_storage, tmp_path, monkeypatch):
    pdf_path = tmp_path / "source.pdf"
    with fitz.open() as doc:
        page = doc.new_page()
        page.insert_text((72, 100), "摘要", fontsize=16, **cjk_font_args())
        page.insert_text(
            (72, 130), "本文研究论文格式检查方法。", fontsize=12, **cjk_font_args()
        )
        doc.save(str(pdf_path))

    key = build_object_key("source.pdf")
    data = pdf_path.read_bytes()
    test_storage.upload(key, io.BytesIO(data), len(data), "application/pdf")
    save_report(db_session, key, _check_report(key, "本文研究论文格式检查方法。"), "check")

    monkeypatch.setattr(
        "app.services.doc_fixer.chat_json",
        lambda messages: {"actions": [{"op": "set_line_spacing", "value": 1.5}]},
    )

    result = fix_document(db_session, key, ["i3"], file_name="PDF修复.docx")
    assert result["converted_from_pdf"] is True
    assert any("pdf2docx" in note for note in result["notes"])
    assert result["applied_actions"]
    Document(io.BytesIO(test_storage.read_bytes(result["file_key"])))  # 可被 Word 打开


# ---------------------------------------------------------------- 空白页规则的边界


def _paragraph_texts(document) -> list[str]:
    return [paragraph.text for paragraph in document.paragraphs]


def test_正文中间的连续空段落不被当成空白页压平():
    """两个连续空段落是作者自己留的段落间距，不是空白页。

    早期实现「连续两个以上空段落即空白页」，会把图表与正文之间、章节之间的留白
    一并删掉，学生拿到的修复稿看起来「更紧凑」，实际是文档被改坏了。
    """
    document = Document()
    document.add_paragraph("正文第一段")
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("正文第二段")

    applied, failed = apply_actions(document, [{"op": "remove_blank_page", "target": "blank_pages"}])

    assert applied == []
    assert failed and failed[0]["reason"] == "没有找到可删除的空白页"
    assert _paragraph_texts(document) == ["正文第一段", "", "", "正文第二段"]


def test_只含软换行的段落不被清除():
    """软换行（Shift+Enter）只是排版换行，不是分页符，也不是空段落。

    把它当空段落删掉，正文里的手工换行会整片消失，而 Word 里看不出区别。
    """
    document = Document()
    document.add_paragraph("正文第一段")
    for _ in range(2):
        soft = document.add_paragraph("")
        soft.add_run().add_break()  # 无 type 的 <w:br/>，即软换行
    document.add_paragraph("参考文献")

    applied, _failed = apply_actions(document, [{"op": "remove_blank_page", "target": "blank_pages"}])

    assert applied == []
    assert len(document.paragraphs) == 4


def test_分页符后的连续空段落作为空白页被清理():
    document = Document()
    document.add_paragraph("第一章")
    document.add_page_break()
    document.add_paragraph("")
    document.add_paragraph("")
    document.add_paragraph("第二章")

    applied, failed = apply_actions(document, [{"op": "remove_blank_page", "target": "blank_pages"}])

    assert failed == []
    assert applied == [{"op": "remove_blank_page", "affected": 2}]
    # 分页符段落保留（删了会把两章并到一页），空白页上的空段落被删
    assert _paragraph_texts(document) == ["第一章", "", "第二章"]


def test_文末多余空段落被清理():
    document = Document()
    document.add_paragraph("正文")
    document.add_paragraph("")
    document.add_paragraph("")

    applied, _failed = apply_actions(document, [{"op": "remove_blank_page", "target": "blank_pages"}])

    assert applied == [{"op": "remove_blank_page", "affected": 2}]
    assert _paragraph_texts(document) == ["正文"]
