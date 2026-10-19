"""服务端生成文件的归属与撤回：AI 修复、PDF 转 Word 的产物必须能撤回。

对象存储不记归属，只有上传记录能回答「这个文件是谁的」。生成文件漏写记录时，
学生点「移除」会被归属校验拒绝：界面上文件消失了，对象却一直留在桶里。
这里从接口层把「生成 → 撤回」整条链路钉住，顺带验证别人撤不走。
"""

from __future__ import annotations

import fitz
import pytest
from docx import Document
from sqlalchemy import select

from app.models.uploaded_file import UploadedFile
from app.services.doc_parser import save_report

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def login(client, device_id: str) -> tuple[dict[str, str], int]:
    """登录并返回 (请求头, 用户 id)。"""
    body = client.post("/api/auth/login", json={"device_id": device_id}).json()
    assert body["code"] == 0, body
    data = body["data"]
    return {"Authorization": f"Bearer {data['token']}"}, int(data["user"]["id"])


def upload(client, headers: dict[str, str], name: str, data: bytes, mime: str) -> str:
    body = client.post(
        "/api/file/upload", files={"file": (name, data, mime)}, headers=headers
    ).json()
    assert body["code"] == 0, body
    return body["data"]["file_key"]


def make_docx(path) -> bytes:
    document = Document()
    document.add_heading("摘要", level=1)
    document.add_paragraph("本文研究论文格式检查方法。")
    document.save(str(path))
    return path.read_bytes()


def make_pdf(path) -> bytes:
    with fitz.open() as document:
        document.new_page()
        document.save(str(path))
    return path.read_bytes()


def check_report(file_key: str) -> dict:
    return {
        "file_key": file_key,
        "issues": [
            {
                "id": "i1",
                "type": "layout",
                "severity": "medium",
                "location": "全文",
                "original": "行距不统一",
                "suggestion": "正文统一 1.5 倍行距",
                "reason": "GB/T 7713.1",
                "source": "rule",
            }
        ],
    }


@pytest.fixture()
def fixed_key(client, db_session, test_storage, tmp_path, monkeypatch):
    """上传一份 docx → 体检 → AI 修复，返回修复副本的 file_key。"""
    headers, owner_id = login(client, "generated-fix-owner")
    key = upload(client, headers, "论文.docx", make_docx(tmp_path / "论文.docx"), DOCX_MIME)
    save_report(db_session, key, check_report(key), "check")
    monkeypatch.setattr(
        "app.services.doc_fixer.chat_json",
        lambda messages: {"actions": [{"op": "set_line_spacing", "target": "body", "value": 1.5}]},
    )

    body = client.post(
        "/api/ai/fix",
        json={"file_key": key, "issue_ids": ["i1"], "file_name": "修复结果.docx"},
        headers=headers,
    ).json()
    assert body["code"] == 0, body
    generated = body["data"]["file_key"]
    assert generated != key
    return headers, owner_id, generated


def test_AI修复生成的文件记在当前用户名下(client, db_session, test_storage, fixed_key):
    _headers, owner_id, generated = fixed_key
    record = db_session.execute(
        select(UploadedFile).where(UploadedFile.file_key == generated)
    ).scalar_one()
    assert record.user_id == owner_id
    assert record.file_name == "修复结果.docx"
    assert record.size > 0


def test_AI修复生成的文件本人可撤回(client, db_session, test_storage, fixed_key):
    headers, _, generated = fixed_key
    assert generated in {obj.key for obj in test_storage.list_objects("files/")}

    body = client.post(
        "/api/file/delete", json={"file_key": generated}, headers=headers
    ).json()
    assert body["code"] == 0, body
    assert body["data"]["deleted"] is True

    # 对象与记录都要消失，不能只让界面上的文件不见
    assert generated not in {obj.key for obj in test_storage.list_objects("files/")}
    assert db_session.execute(
        select(UploadedFile).where(UploadedFile.file_key == generated)
    ).scalar_one_or_none() is None


def test_AI修复生成的文件别人撤不走(client, db_session, test_storage, fixed_key):
    _, _, generated = fixed_key
    other, _ = login(client, "generated-fix-other")

    body = client.post(
        "/api/file/delete", json={"file_key": generated}, headers=other
    ).json()
    assert body["code"] == 40404
    # 记录还在，本人仍然撤得走
    assert db_session.execute(
        select(UploadedFile).where(UploadedFile.file_key == generated)
    ).scalar_one_or_none() is not None
    assert generated in {obj.key for obj in test_storage.list_objects("files/")}


def test_PDF转Word生成的文件可撤回(client, db_session, test_storage, tmp_path):
    headers, owner_id = login(client, "generated-convert-owner")
    pdf_key = upload(client, headers, "扫描件.pdf", make_pdf(tmp_path / "扫描件.pdf"), "application/pdf")

    body = client.post(
        "/api/ai/convert",
        json={"file_key": pdf_key, "target": "docx", "file_name": "转换结果.docx"},
        headers=headers,
    ).json()
    assert body["code"] == 0, body
    generated = body["data"]["file_key"]

    record = db_session.execute(
        select(UploadedFile).where(UploadedFile.file_key == generated)
    ).scalar_one()
    assert record.user_id == owner_id

    other, _ = login(client, "generated-convert-other")
    rejected = client.post(
        "/api/file/delete", json={"file_key": generated}, headers=other
    ).json()
    assert rejected["code"] == 40404

    deleted = client.post(
        "/api/file/delete", json={"file_key": generated}, headers=headers
    ).json()
    assert deleted["code"] == 0, deleted
    assert generated not in {obj.key for obj in test_storage.list_objects("files/")}
    assert db_session.execute(
        select(UploadedFile).where(UploadedFile.file_key == generated)
    ).scalar_one_or_none() is None
