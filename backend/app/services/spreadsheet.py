"""Excel 归一化：把「页数取决于排版」的表格变成「页数确定」。

同一个汇总表，页数会随四件事变化：列宽与缩放决定横向切几刀、默认只打印活动工作表、
打印区域与手动分页符、隐藏行列与空白区造出的空页。学生看到的「1 页」和实际出纸的
「12 页」一旦不一致，收银台就要吵价格。

所以转换前统一版式，让页数只由内容决定：

- 缩到一页宽，纵向不限（宽表不会横向切碎）
- 打印区域收紧到真正有内容的行列，去掉尾部空行空列
- 清掉手动分页符与重复标题行，避免额外分页
- 只打可见工作表；隐藏表是学生主动藏起来的，不该出纸

归一化时同时把「有效区域之外还有多少空行空列」与最终出纸的中位字号带出来，
供体检规则判断。这里只产出事实，不判断是否报问题。

改口径只改 config/spreadsheet_rules.json。
"""

from __future__ import annotations

import logging
import time
from contextlib import contextmanager
from contextvars import ContextVar
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import ColBreak, RowBreak
from openpyxl.worksheet.properties import PageSetupProperties

from app.config import get_spreadsheet_rules, settings
from app.utils.exceptions import BusinessError

logger = logging.getLogger(__name__)

# 一条表格请求允许花在「归一版式 + 转换」上的总时长（见 settings 里的同名配置）。
# 每放宽一次横向页数，就要重跑一轮归一化 + 转换，期间全程占住一个 AI 并发名额。
#
# 光限次数不够，只限总时间也不够，两个都是硬约束：
# - 次数：横向页数 1..max_width_pages（默认 4），到上限就按当前结果出纸；
# - 时间：第一轮给满一次转换超时（不给连一份 PDF 都拿不出来），之后每一轮只能用
#   「预算 - 已用」，剩余时间用尽就不再重试。于是最坏总耗时 = max(预算, 单次转换超时)，
#   默认 240 秒。
#
# 为什么两轮都要把剩余时间交给转换层（见 render_deadline）：验收实测过「每次转换
# 130 秒时仍跑 2 次、占用 260 秒」—— 因为此前只看「已用 ≥ 预算」才决定停不停，
# 第二轮一旦开始就能跑满自己的 180 秒，预算形同虚设。真正吃容量的是时间，所以
# 额度必须传到真正等待 soffice 的那一层，预算耗尽即杀进程、归还名额。
# 超预算就按已有结果出纸，并在体检报告里写明「因超时按前一次结果出纸」，
# 让用户自己拆表 —— 不拿全站容量去碰运气。
RENDER_BUDGET_SECONDS_FALLBACK = 240.0
# 单次转换的默认上限（settings.soffice_timeout_seconds 的兜底值）
SINGLE_CONVERSION_SECONDS_FALLBACK = 180.0
RENDER_TIMEOUT_NOTE = "因超时按前一次结果出纸"
RENDER_FAILED_NOTE = "放宽横向页数时的转换失败，按前一次结果出纸"

# 当前这次转换的截止时刻（monotonic）。用 ContextVar 而不是参数透传：转换函数
# 在 doc_parser 里，签名被十几个调用方与测试替身共用，为了一次预算就改契约
# 得不偿失；而 render_pdf 与转换在同一线程同一调用栈里跑，传递范围刚好够用。
_render_deadline: ContextVar[float | None] = ContextVar(
    "spreadsheet_render_deadline", default=None
)


@contextmanager
def render_deadline(seconds: float | None):
    """在 with 里跑的转换最多只能等 seconds 秒；None 表示不受表格预算约束。"""
    if seconds is None:
        yield
        return
    token = _render_deadline.set(time.monotonic() + max(0.0, seconds))
    try:
        yield
    finally:
        _render_deadline.reset(token)


def remaining_render_seconds() -> float | None:
    """这次转换还剩多少秒可用；不在表格渲染里（Word/PDF 转换）时返回 None。"""
    deadline = _render_deadline.get()
    if deadline is None:
        return None
    return max(0.0, deadline - time.monotonic())


def _is_blank(value: object) -> bool:
    if value is None:
        return True
    return isinstance(value, str) and not value.strip()


def _content_bounds(worksheet) -> tuple[int, int, int, int] | None:
    """真正有内容的行列范围；整张表都空则返回 None。"""
    max_row = worksheet.max_row or 0
    while max_row > 0 and all(_is_blank(c.value) for c in worksheet[max_row]):
        max_row -= 1
    if max_row == 0:
        return None

    last_col = 0
    for row in worksheet.iter_rows(max_row=max_row):
        for cell in row:
            if not _is_blank(cell.value) and cell.column > last_col:
                last_col = cell.column
    if last_col == 0:
        return None
    return 1, 1, max_row, last_col


def _apply_page_setup(
    worksheet, bounds: tuple[int, int, int, int], rules: dict, fit_to_width: int
) -> None:
    from openpyxl.utils import get_column_letter

    first_row, first_col, last_row, last_col = bounds
    worksheet.print_area = f"{get_column_letter(first_col)}{first_row}:{get_column_letter(last_col)}{last_row}"

    worksheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    worksheet.page_setup.fitToWidth = fit_to_width
    worksheet.page_setup.fitToHeight = 0  # 纵向不限，按内容往下分页
    worksheet.page_setup.orientation = str(rules.get("orientation", "landscape"))
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4

    margin_cm = float(rules.get("margin_cm", 1.0))
    margin_inch = margin_cm / 2.54
    worksheet.page_margins = PageMargins(
        left=margin_inch,
        right=margin_inch,
        top=margin_inch,
        bottom=margin_inch,
        header=0.0,
        footer=0.0,
    )
    worksheet.print_options.horizontalCentered = True

    # 手动分页符会把一张表切碎，且学生自己也不一定记得设过
    worksheet.row_breaks = RowBreak()
    worksheet.col_breaks = ColBreak()
    worksheet.print_title_rows = None
    # 横向分成多页时把第一列当标签列重复，否则第二页起不知道每行是谁
    worksheet.print_title_cols = "A:A" if fit_to_width > 1 else None


def normalize_spreadsheet(
    source: Path,
    work_dir: Path,
    fit_to_width: int = 1,
    stats: dict | None = None,
) -> Path:
    """按统一版式重存一份表格，返回新文件路径。原文件不动。

    fit_to_width 是横向占用几页：默认 1（缩到一页宽），列太多看不清时由调用方放宽。
    stats 不为 None 时，把每个可见工作表的有效区域与末尾空行空列写进去。
    """
    rules = get_spreadsheet_rules()
    try:
        workbook = load_workbook(source, data_only=False)
    except Exception as exc:
        raise BusinessError("表格文件打不开，可能已损坏或不是 Excel 格式", code=50043) from exc

    printable = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
    if not printable:
        raise BusinessError("表格里没有可见的工作表，请先取消隐藏再上传", code=50044)

    kept = 0
    sheet_stats: list[dict] = []
    for worksheet in printable:
        bounds = _content_bounds(worksheet)
        if bounds is None:
            worksheet.sheet_state = "hidden"  # 空表不出纸，也不占页
            logger.info("工作表「%s」没有任何内容，跳过", worksheet.title)
            continue
        _apply_page_setup(worksheet, bounds, rules, fit_to_width)
        kept += 1
        if stats is not None:
            first_row, first_col, last_row, last_col = bounds
            sheet_stats.append(
                {
                    "title": worksheet.title,
                    "content_rows": last_row - first_row + 1,
                    "content_columns": last_col - first_col + 1,
                    # openpyxl 的 max_row/max_column 会被格式或幽灵单元格撑大，
                    # 与有效区域相减就是学生自己也不一定看得见的尾部空白
                    "trailing_blank_rows": max(0, int(worksheet.max_row or 0) - last_row),
                    "trailing_blank_cols": max(0, int(worksheet.max_column or 0) - last_col),
                }
            )

    if kept == 0:
        raise BusinessError("表格里没有任何内容，请检查后再上传", code=50045)

    if stats is not None:
        stats["sheets"] = sheet_stats

    work_dir.mkdir(parents=True, exist_ok=True)
    target = work_dir / f"{source.stem}.normalized.xlsx"
    workbook.save(target)
    logger.info(
        "表格已归一版式：%s，打印 %d 个工作表，跳过隐藏 %d 个",
        source.name,
        kept,
        len(workbook.worksheets) - kept,
    )
    return target


def _print_font_percentile(pdf_path: Path) -> float | None:
    """转换结果里的中位字号。用中位数而不是最小值：少数脚注不该绑架整表版式。"""
    sizes: list[float] = []
    document = pymupdf.open(pdf_path)
    try:
        for page in document:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        if span["text"].strip():
                            sizes.append(float(span["size"]))
    finally:
        document.close()
    if not sizes:
        return None
    sizes.sort()
    return sizes[len(sizes) // 2]


def render_pdf(source: Path, work_dir: Path, convert, stats: dict | None = None) -> Path:
    """把表格渲染成 PDF，并保证字不至于小到看不清。

    缩到一页宽是最省纸也最合直觉的做法，但列一多，缩放会把字压到五六号，
    出纸等于白印。所以先按一页宽试，中位字号低于阈值就把横向页数放宽一页重来，
    直到字够大或到上限。首列会作为标签列重复，横向分页也不会丢行头。

    次数上限（max_width_pages）与总耗时上限（spreadsheet_render_budget_seconds）
    都是硬约束：每一轮的等待额度交给转换层，预算用尽即停，按已有的那份出纸。

    stats 不为 None 时，把最终实际使用的横向页数、中位字号、各表空白区，
    以及是否因超时降级（render_truncated / render_note / 已用时间与次数）写进去；
    体检规则据此判断「列太多字号过小」与「末尾大片空行空列」，
    返回体里的 render_note 则直接告诉用户为什么没有继续放宽。
    """
    rules = get_spreadsheet_rules()
    min_pt = float(rules.get("min_readable_font_pt", 8.0))
    max_width = max(1, int(rules.get("max_width_pages", 4)))
    budget = max(0.0, float(getattr(settings, "spreadsheet_render_budget_seconds", RENDER_BUDGET_SECONDS_FALLBACK)))
    single = max(
        1.0, float(getattr(settings, "soffice_timeout_seconds", SINGLE_CONVERSION_SECONDS_FALLBACK))
    )
    started = time.monotonic()

    pdf_path: Path | None = None
    attempt_stats: dict = {}
    attempts = 0
    note = ""

    for width_pages in range(1, max_width + 1):
        if width_pages == 1:
            # 第一轮必须跑满一次转换的时间：没有它连一份 PDF 都交不出来，
            # 预算再紧也不能把它卡成半份
            cap = max(budget, single)
        else:
            cap = budget - (time.monotonic() - started)
            if cap <= 0:
                logger.warning(
                    "表格渲染已用掉 %.0f 秒预算（上限 %.0f 秒），不再放宽横向页数：%s",
                    time.monotonic() - started,
                    budget,
                    source.name,
                )
                note = RENDER_TIMEOUT_NOTE
                break

        attempt_dir = work_dir / f"fit{width_pages}"
        attempt_dir.mkdir(parents=True, exist_ok=True)
        pending_stats: dict = {}
        try:
            # 归一化也算在这一轮的额度里：大表 load_workbook + save 本身就不便宜
            with render_deadline(cap):
                prepared = normalize_spreadsheet(
                    source, attempt_dir, fit_to_width=width_pages, stats=pending_stats
                )
                produced = convert(prepared, attempt_dir)
        except BusinessError as exc:
            if pdf_path is None:
                raise  # 第一轮就失败：手里没有任何结果，照旧报错
            logger.warning(
                "表格横向 %d 页重试失败（%s），按前一次结果出纸：%s",
                width_pages,
                exc.message,
                source.name,
            )
            # 额度比单次转换超时更紧，说明是被预算卡死的；否则是转换自己失败
            note = RENDER_TIMEOUT_NOTE if cap < single else RENDER_FAILED_NOTE
            break

        attempts = width_pages
        pdf_path = produced
        attempt_stats = pending_stats
        font_pt = _print_font_percentile(pdf_path)
        attempt_stats.update(
            {
                "fit_to_width_pages": width_pages,
                "print_font_pt": font_pt,
                "min_readable_font_pt": min_pt,
            }
        )
        if font_pt is None or font_pt >= min_pt:
            logger.info(
                "表格渲染完成：%s，横向 %d 页，中位字号 %.1f pt",
                source.name,
                width_pages,
                font_pt or 0.0,
            )
            break
        logger.info(
            "表格缩到横向 %d 页后中位字号只有 %.1f pt，低于 %.1f pt，放宽一页重试",
            width_pages,
            font_pt,
            min_pt,
        )
    else:
        logger.warning("表格横向已放宽到 %d 页仍偏小，按最大宽度出纸：%s", max_width, source.name)

    if pdf_path is None:  # 只有 max_width < 1 这种配置错位才可能走到，兜底不返回 None
        raise BusinessError("表格渲染失败，请稍后重试", code=50047)

    attempt_stats["render_attempts"] = attempts
    attempt_stats["render_elapsed_seconds"] = round(time.monotonic() - started, 2)
    attempt_stats["render_budget_seconds"] = budget
    if note:
        attempt_stats["render_truncated"] = True
        attempt_stats["render_note"] = note
    if stats is not None:
        stats.update(attempt_stats)
    return pdf_path
