"""生成答辩演示用的「含格式错误的毕业论文」docx。

用法：
    cd backend
    .venv\\Scripts\\python.exe scripts\\make_demo_docx.py [输出路径]

刻意埋入的问题（对应交接文档附录的演示道具清单）：
1. 缺少「目录」——结构缺失
2. 左边距 2.0cm（规则要求 3.0cm）——页边距偏差
3. 一张 19cm 宽的图片——超出 A4 版心
4. 一个空白页——浪费纸张
5. 参考文献著录错误：缺页码、作者超三人未用「等」、缺出版年
6. 若干错别字与标点问题（供阶段 5-C 的错别字检查使用）
"""

from __future__ import annotations

import sys
from pathlib import Path

import fitz
from docx import Document
from docx.shared import Cm

DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "runtime" / "demo_thesis_flawed.docx"


def _make_chart(path: Path) -> Path:
    """生成一张占位图，用于演示「图片超出页边距」。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 800, 400))
    pix.clear_with(220)
    pix.save(str(path))
    return path


def build(output: Path) -> Path:
    output.parent.mkdir(parents=True, exist_ok=True)
    chart = _make_chart(output.parent / "demo_chart.png")

    doc = Document()
    section = doc.sections[0]
    section.left_margin = Cm(2.0)   # 规则要求 3.0cm，故意偏差
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    doc.add_heading("摘要", level=1)
    doc.add_paragraph(
        "本文围绕高校学位论文的格式检查问题,提出了一种规则引擎与大模型混合的检查方法。"
        "实验结杲表明,该方法能够有效识别论文中的格式问题。"
    )
    doc.add_paragraph("关键词：论文格式；自动检查；大模型；文印室")

    doc.add_page_break()
    doc.add_heading("第一章 绪论", level=1)
    doc.add_paragraph("随着高校学位论文数量逐年增加，格式审查工作给文印室带来了很大压力。" * 3)
    doc.add_paragraph("本文提出的方法包含以下特点：")
    doc.add_picture(str(chart), width=Cm(19))  # 超出 A4 版心宽度

    doc.add_page_break()
    doc.add_page_break()  # 制造一个空白页

    doc.add_heading("第二章 系统设计", level=1)
    doc.add_paragraph("系统分为学生端与老师端两部分，采用前后端分离架构。")

    doc.add_page_break()
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] 张三. 论文格式自动检查方法研究[J]. 计算机应用, 2020, 12(3).")
    doc.add_paragraph(
        "[2] 李四, 王五, 赵六, 孙七. 大模型在文档结构化处理中的应用[M]. 北京: 科学出版社."
    )
    doc.add_paragraph("[3] 王小明. 高校文印服务优化研究[D]. 上海: 某大学, 2021.")

    doc.save(str(output))
    return output


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUTPUT
    saved = build(target)
    print(f"演示文档已生成：{saved}")
